import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# -----------------------------------------------------------------------------
# Color Palette & Design Constants
# -----------------------------------------------------------------------------
BG_COLOR = RGBColor(15, 23, 42)          # Deep Slate (#0F172A)
CARD_BG = RGBColor(30, 41, 59)           # Dark Slate Navy (#1E293B)
CARD_BORDER = RGBColor(51, 65, 85)       # Border Slate (#334155)
CARD_BG_ALT = RGBColor(24, 32, 47)       # Slightly darker card (#18202F)

TEXT_LIGHT = RGBColor(248, 250, 252)     # Off-white (#F8FAFC)
TEXT_MUTED = RGBColor(148, 163, 184)     # Slate Grey (#94A3B8)
TEXT_DIM = RGBColor(100, 116, 139)       # Dim Grey (#64748B)

ACCENT_TERRACOTTA = RGBColor(224, 86, 76) # SIES Crimson/Terracotta (#E0564C)
ACCENT_GOLD = RGBColor(229, 184, 105)    # Antique Gold (#E5B869)
ACCENT_CYAN = RGBColor(56, 189, 248)     # Tech Cyan (#38BDF8)
ACCENT_EMERALD = RGBColor(52, 211, 153)  # Muted Jade/Emerald (#34D399)
ACCENT_PURPLE = RGBColor(192, 132, 252)  # Soft Violet (#C084FC)

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

def add_header(slide, category, title, subtitle):
    set_slide_background(slide)
    
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p0 = tf.paragraphs[0]
    p0.text = category.upper()
    p0.font.name = FONT_HEADING
    p0.font.size = Pt(9.5)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_GOLD
    p0.space_after = Pt(2)
    
    p1 = tf.add_paragraph()
    p1.text = title
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(20)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT
    p1.space_after = Pt(2)
    
    p2 = tf.add_paragraph()
    p2.text = subtitle
    p2.font.name = FONT_BODY
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_MUTED

def add_card(slide, left, top, width, height, border_color=CARD_BORDER, bg_color=CARD_BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.2)
    else:
        shape.line.fill.background()
    return shape

# -----------------------------------------------------------------------------
# Slide Generators
# -----------------------------------------------------------------------------

def build_slide_1(prs):
    # Slide 1: Title Slide (Cover)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.3), Inches(7.5))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_TERRACOTTA
    bar.line.fill.background()
    
    add_card(slide, Inches(1.0), Inches(0.8), Inches(11.333), Inches(5.9), border_color=ACCENT_GOLD, bg_color=CARD_BG)
    
    logo_path = "logo.png"
    if os.path.isfile(logo_path):
        slide.shapes.add_picture(logo_path, Inches(1.6), Inches(1.3), height=Inches(1.5))
    
    tb = slide.shapes.add_textbox(Inches(3.3), Inches(1.2), Inches(8.5), Inches(3.2))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = "SIES (DR. APJ ABDUL KALAM MEMORIAL) HIGH SCHOOL"
    p0.font.name = FONT_HEADING
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_GOLD
    p0.space_after = Pt(6)
    
    p1 = tf.add_paragraph()
    p1.text = "Academic Timetable Generator"
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(30)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT
    p1.space_after = Pt(4)
    
    p2 = tf.add_paragraph()
    p2.text = "High-Performance Multi-Constraint School Scheduling & Desktop Management Platform"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(13)
    p2.font.color.rgb = ACCENT_CYAN
    p2.space_after = Pt(16)
    
    p3 = tf.add_paragraph()
    p3.text = "Autonomous Two-Phase CP-SAT Optimization  *  Independent Zero-Trust Verification  *  Tauri v2 Desktop Shell"
    p3.font.name = FONT_BODY
    p3.font.size = Pt(11)
    p3.font.color.rgb = TEXT_MUTED

    meta_box = slide.shapes.add_textbox(Inches(1.6), Inches(4.5), Inches(10.1), Inches(1.8))
    mtf = meta_box.text_frame
    mtf.word_wrap = True
    
    mp0 = mtf.paragraphs[0]
    mp0.text = "PROJECT METADATA & AUTHORSHIP"
    mp0.font.name = FONT_HEADING
    mp0.font.size = Pt(9.5)
    mp0.font.bold = True
    mp0.font.color.rgb = ACCENT_GOLD
    mp0.space_after = Pt(4)
    
    mp1 = mtf.add_paragraph()
    mp1.text = "Author / Lead Developer: Sourish Kumar Rai   |   Software Engineering Capstone Showcase"
    mp1.font.name = FONT_BODY
    mp1.font.size = Pt(12)
    mp1.font.bold = True
    mp1.font.color.rgb = TEXT_LIGHT
    mp1.space_after = Pt(4)
    
    mp2 = mtf.add_paragraph()
    mp2.text = "Core Stack: Java 21 LTS  *  Google OR-Tools CP-SAT 9.11  *  React 19 + TypeScript  *  Tauri v2 + Rust"
    mp2.font.name = FONT_BODY
    mp2.font.size = Pt(10.5)
    mp2.font.color.rgb = ACCENT_EMERALD

def build_slide_2(prs):
    # Slide 2: Problem Statement & Academic Context
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Problem Definition & Context", "The Academic Timetabling Challenge",
               "Why Secondary School Scheduling is an NP-Hard Combinatorial Bottleneck")
    
    col_w = Inches(3.64)
    gap = Inches(0.4)
    top_pos = Inches(1.7)
    card_h = Inches(5.1)
    
    # Card 1: Combinatorial Scale
    add_card(slide, Inches(0.8), top_pos, col_w, card_h)
    tb = slide.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Combinatorial Scale"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TERRACOTTA
    p.space_after = Pt(10)
    
    bullets = [
        ("Massive Decision Space: ", "26 classes (Standards V to X), 44 faculty members, and 25 subjects across 6 days."),
        ("1,248 Timetable Slots: ", "Every single weekly period must be populated without gaps or double bookings."),
        ("Exponential Complexity: ", "The School Timetabling Problem (STP) is formally NP-Complete. Naive search space exceeds 10^300 states."),
        ("Dynamic Section Sizes: ", "Standards V & VIII have 5 divisions; VI, VII, IX & X have 4 divisions.")
    ]
    for b_title, b_desc in bullets:
        bp = tf.add_paragraph()
        bp.text = "* " + b_title + b_desc
        bp.font.name = FONT_BODY
        bp.font.size = Pt(10.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(8)

    # Card 2: Manual Failure Modes
    add_card(slide, Inches(0.8) + col_w + gap, top_pos, col_w, card_h)
    tb = slide.shapes.add_textbox(Inches(1.0) + col_w + gap, top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Manual Failure Modes"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD
    p.space_after = Pt(10)
    
    bullets = [
        ("Teacher Clashes: ", "Manual spreadsheets inevitably lead to simultaneous overlapping assignments for faculty."),
        ("Consecutive Burnout: ", "Teachers assigned 4-5 continuous periods without recovery breaks or recess spacing."),
        ("Shared Resource Contention: ", "Multiple classes scheduled into the sole Computer Lab during the same period."),
        ("Upstream Discrepancies: ", "Standard-wise curriculum demands frequently mismatch teacher-wise workload sheets.")
    ]
    for b_title, b_desc in bullets:
        bp = tf.add_paragraph()
        bp.text = "* " + b_title + b_desc
        bp.font.name = FONT_BODY
        bp.font.size = Pt(10.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(8)

    # Card 3: The Solution Imperative
    add_card(slide, Inches(0.8) + (col_w + gap) * 2, top_pos, col_w, card_h)
    tb = slide.shapes.add_textbox(Inches(1.0) + (col_w + gap) * 2, top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "The Solution Imperative"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.space_after = Pt(10)
    
    bullets = [
        ("Autonomous Generation: ", "Eliminate 3-4 weeks of arduous manual committee scheduling each academic term."),
        ("Mathematical Proof: ", "Replace guesswork with constraint programming proving zero hard constraint violations."),
        ("Pedagogical Optimization: ", "Evenly distribute subject fatigue and balance teacher workloads across weekdays."),
        ("Local Desktop Autonomy: ", "Keep institutional scheduling confidential and offline on local school PCs.")
    ]
    for b_title, b_desc in bullets:
        bp = tf.add_paragraph()
        bp.text = "* " + b_title + b_desc
        bp.font.name = FONT_BODY
        bp.font.size = Pt(10.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(8)

def build_slide_3(prs):
    # Slide 3: Project Objectives & Scope
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "System Goals & Scope", "Project Objectives & Design Mandates",
               "Engineering a Mathematically Sound, Turnkey Desktop Scheduling Solution")
    
    w = Inches(5.6)
    h = Inches(2.4)
    x1 = Inches(0.8)
    x2 = Inches(6.8)
    y1 = Inches(1.7)
    y2 = Inches(4.4)
    
    quads = [
        (x1, y1, "1. Zero-Defect Hard Constraint Solvability", ACCENT_TERRACOTTA, [
            "Enforce 10 non-negotiable operational rules (H1 to H10) strictly.",
            "Guarantee 0 teacher double-bookings, 0 computer lab collisions, and 100% allotment fulfillment.",
            "Independent post-solve audit evaluating every slot from scratch without relying on solver state."
        ]),
        (x2, y1, "2. Multi-Objective Pedagogical Optimization", ACCENT_GOLD, [
            "Formulate and minimize weighted soft penalties across 6 objectives (S1 to S6).",
            "Uniformly space multi-period subjects across weekdays to prevent student cognitive fatigue.",
            "Balance faculty daily lecture counts and eliminate wasted idle waiting windows."
        ]),
        (x1, y2, "3. Turnkey Windows Desktop Packaging", ACCENT_CYAN, [
            "Package as a standalone Windows desktop app via Tauri v2 and NSIS (58.5 MB installer).",
            "Bundle a private 48 MB JRE (jlink) with pre-staged Google OR-Tools native Windows binaries.",
            "Zero prerequisites for school administrators: no Java, Node, or terminal interaction required."
        ]),
        (x2, y2, "4. Interactive Timetable Studio & In-Place Editing", ACCENT_EMERALD, [
            "Modern React 19 UI with SIES school branding (Terracotta, Crimson, Gold palette).",
            "Interactive Day * Period weekly grids for all 26 classes and 44 teachers.",
            "In-place slot swapping with live constraint validation and full Undo/Redo history."
        ])
    ]
    
    for x, y, title, color, points in quads:
        add_card(slide, x, y, w, h)
        tb = slide.shapes.add_textbox(x + Inches(0.25), y + Inches(0.15), w - Inches(0.5), h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_HEADING
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color
        p.space_after = Pt(6)
        
        for pt in points:
            bp = tf.add_paragraph()
            bp.text = "* " + pt
            bp.font.name = FONT_BODY
            bp.font.size = Pt(10)
            bp.font.color.rgb = TEXT_LIGHT
            bp.space_after = Pt(4)

def build_slide_4(prs):
    # Slide 4: Technology Stack & Justification
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Technical Foundations", "Technology Stack & Architecture Justification",
               "Purpose-Built Multi-Tier Architecture Engineered for Performance and Reliability")
    
    y_start = Inches(1.7)
    card_h = Inches(1.15)
    gap = Inches(0.18)
    card_w = Inches(11.733)
    
    layers = [
        ("CORE SOLVER TIER", "Java 21 LTS + Google OR-Tools CP-SAT 9.11", ACCENT_TERRACOTTA,
         "Selected for industrial-grade constraint satisfaction. CP-SAT applies Conflict-Driven Clause Learning (CDCL) and lazy clause generation to explore 18,800+ variables with multi-threaded parallel search. Vastly outperforms slow genetic algorithms."),
        ("API DAEMON TIER", "Lightweight Embedded Java HTTP Server (jdk.httpserver)", ACCENT_GOLD,
         "Zero external web framework overhead (no Spring/Tomcat bloat). Strictly bound to local loopback 127.0.0.1:8080 with dynamic ephemeral port fallback. Exposes clean REST endpoints for status, solving, studio grids, and validation."),
        ("DESKTOP SHELL TIER", "Tauri v2 + Rust Native Windows Shell", ACCENT_CYAN,
         "Provides a native Windows application wrapper utilizing system WebView2. Tiny 8.9 MB desktop binary vs 150+ MB memory-heavy Electron shells. Native Win32 Job Object manages the child Java process to guarantee zero orphaned background tasks."),
        ("PRESENTATION & EXPORT", "React 19 + TypeScript + Vite | Apache POI & OpenPDF", ACCENT_EMERALD,
         "Modern responsive frontend using custom CSS design tokens reflecting SIES identity. High-fidelity publishing engine generating multi-tab Excel workbooks (.xlsx), 26 per-class vector PDFs, and 9 machine-readable audit CSV reports.")
    ]
    
    for idx, (tag, tech, color, rationale) in enumerate(layers):
        top_pos = y_start + idx * (card_h + gap)
        add_card(slide, Inches(0.8), top_pos, card_w, card_h)
        
        tb_left = slide.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.15), Inches(3.6), card_h - Inches(0.3))
        tf_left = tb_left.text_frame
        tf_left.word_wrap = True
        
        p0 = tf_left.paragraphs[0]
        p0.text = tag
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(9.5)
        p0.font.bold = True
        p0.font.color.rgb = color
        p0.space_after = Pt(2)
        
        p1 = tf_left.add_paragraph()
        p1.text = tech
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_LIGHT
        
        tb_right = slide.shapes.add_textbox(Inches(4.8), top_pos + Inches(0.15), Inches(7.5), card_h - Inches(0.3))
        tf_right = tb_right.text_frame
        tf_right.word_wrap = True
        p_desc = tf_right.paragraphs[0]
        p_desc.text = rationale
        p_desc.font.name = FONT_BODY
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = TEXT_MUTED

def build_slide_5(prs):
    # Slide 5: System Architecture & Component Interaction
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "System Architecture", "Desktop Architecture & Process Lifecycle",
               "Strictly Isolated Client-Daemon Communication with Guaranteed Process Teardown")
    
    box_w = Inches(3.6)
    box_h = Inches(4.8)
    top_pos = Inches(1.8)
    gap = Inches(0.46)
    
    # Box 1: Tauri Frontend
    add_card(slide, Inches(0.8), top_pos, box_w, box_h, border_color=ACCENT_CYAN)
    tb = slide.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.2), box_w - Inches(0.4), box_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "1. Native Desktop Shell"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.space_after = Pt(8)
    
    items = [
        ("Tauri v2 + Rust Core: ", "Launches app.exe, registers window events, manages state."),
        ("WebView2 Frontend: ", "React 19 + TypeScript rendering Command Center, Setup & Studio."),
        ("Headless Child Spawn: ", "Spawns private java.exe using Win32 CREATE_NO_WINDOW (zero popups)."),
        ("Win32 Job Object: ", "Assigns Java PID to JOB_OBJECT_LIMIT_KILL_ON_JOB_CLOSE. Automatically terminates Java on app crash/close.")
    ]
    for h_txt, b_txt in items:
        bp = tf.add_paragraph()
        bp.text = "* " + h_txt + b_txt
        bp.font.name = FONT_BODY
        bp.font.size = Pt(9.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(6)

    # Box 2: IPC & Local Daemon
    add_card(slide, Inches(0.8) + box_w + gap, top_pos, box_w, box_h, border_color=ACCENT_GOLD)
    tb = slide.shapes.add_textbox(Inches(1.0) + box_w + gap, top_pos + Inches(0.2), box_w - Inches(0.4), box_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "2. Local Loopback API"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD
    p.space_after = Pt(8)
    
    items = [
        ("Dynamic Port Hunting: ", "Tests 127.0.0.1:8080. If busy, binds ephemeral port dynamically."),
        ("Health Check Handshake: ", "Tauri polls /api/health until HTTP 200 OK before initializing UI."),
        ("Strict Loopback Binding: ", "Restricted to 127.0.0.1 loopback; rejects external network traffic."),
        ("REST Endpoint Suite: ", "/api/solve/start, /status, /schedule/class, /edit/validate, /export.")
    ]
    for h_txt, b_txt in items:
        bp = tf.add_paragraph()
        bp.text = "* " + h_txt + b_txt
        bp.font.name = FONT_BODY
        bp.font.size = Pt(9.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(6)

    # Box 3: Authoritative Java Engine
    add_card(slide, Inches(0.8) + (box_w + gap) * 2, top_pos, box_w, box_h, border_color=ACCENT_TERRACOTTA)
    tb = slide.shapes.add_textbox(Inches(1.0) + (box_w + gap) * 2, top_pos + Inches(0.2), box_w - Inches(0.4), box_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "3. Authoritative Core"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TERRACOTTA
    p.space_after = Pt(8)
    
    items = [
        ("Private Bundled JRE: ", "Stripped 48 MB JRE (jlink) with jdk.zipfs & jdk.httpserver."),
        ("Google OR-Tools CP-SAT: ", "Two-phase mathematical solver generating 1,248 conflict-free slots."),
        ("ScheduleValidator: ", "Independent zero-trust engine auditing every constraint from scratch."),
        ("Document Exporters: ", "Apache POI master Excel & OpenPDF vector class timetable generation.")
    ]
    for h_txt, b_txt in items:
        bp = tf.add_paragraph()
        bp.text = "* " + h_txt + b_txt
        bp.font.name = FONT_BODY
        bp.font.size = Pt(9.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(6)

def build_slide_6(prs):
    # Slide 6: End-to-End Timetable Generation Workflow
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Execution Pipeline", "End-to-End Timetable Generation Workflow",
               "A Continuous 6-Stage Pipeline Transforming Input Spreadsheets into Validated Schedules")
    
    card_w = Inches(3.64)
    card_h = Inches(2.35)
    gap_x = Inches(0.4)
    gap_y = Inches(0.3)
    x0 = Inches(0.8)
    y0 = Inches(1.8)
    
    steps = [
        ("STAGE 1: Ingestion", "Dual-Source Parsing", ACCENT_CYAN,
         "Parses standard-wise and teacher-wise allocation spreadsheets (.xlsx) alongside class-teacher mappings and constraint definitions using Apache POI."),
        ("STAGE 2: Reconciliation", "Data Normalization & ETL", ACCENT_GOLD,
         "Cross-checks standard demand against teacher supply at the atomic period level. Resolves teacher aliases via NormalizationDictionary; logs anomalies."),
        ("STAGE 3: Pre-Solve Audit", "Invariant Axiom Checking", ACCENT_TERRACOTTA,
         "Verifies foundational scheduling invariants (48 periods/week per class, total 1,248 periods) before invoking solver to prevent unsolvable runs."),
        ("STAGE 4: CP-SAT Solve", "Two-Phase Optimization", ACCENT_EMERALD,
         "Phase 1 solves pure hard constraint feasibility (H1 to H10 in ~4s). Phase 2 warm-starts soft constraint optimization (S1 to S6) minimizing penalties under time limit."),
        ("STAGE 5: Verification", "Zero-Trust Independent Audit", ACCENT_PURPLE,
         "ScheduleValidator audits all 1,248 slots independently of solver internal state, certifying 0 hard violations and computing the objective quality score."),
        ("STAGE 6: Publishing", "Multi-Format Document Export", ACCENT_CYAN,
         "Generates master multi-tab Excel workbook (timetable.xlsx), 26 printable per-class vector PDFs, and 9 machine-readable audit CSV reports.")
    ]
    
    for idx, (stage, title, color, desc) in enumerate(steps):
        row = idx // 3
        col = idx % 3
        cx = x0 + col * (card_w + gap_x)
        cy = y0 + row * (card_h + gap_y)
        
        add_card(slide, cx, cy, card_w, card_h)
        tb = slide.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.15), card_w - Inches(0.4), card_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = stage
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(9.5)
        p0.font.bold = True
        p0.font.color.rgb = color
        p0.space_after = Pt(2)
        
        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_LIGHT
        p1.space_after = Pt(4)
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = TEXT_MUTED

def build_slide_7(prs):
    # Slide 7: Constraint Programming Formulation (OR-Tools CP-SAT)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Mathematical Modeling", "Constraint Formulation with OR-Tools CP-SAT",
               "Mapping Institutional Policies to SAT Decision Variables and Linear Expressions")
    
    col_w = Inches(5.65)
    gap = Inches(0.43)
    top_pos = Inches(1.7)
    card_h = Inches(5.1)
    
    # Col 1: Mathematical Decision Model
    add_card(slide, Inches(0.8), top_pos, col_w, card_h)
    tb = slide.shapes.add_textbox(Inches(1.05), top_pos + Inches(0.2), col_w - Inches(0.5), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Mathematical Decision Model"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.space_after = Pt(8)
    
    pts = [
        ("Boolean Decision Variables: ", "x(c, d, p, a) in {0, 1} indicates whether atomic allocation a = (c, s, t) is scheduled in class c on day d at period p."),
        ("Model Scale (Production 26 Classes): ", "Creates exactly 18,864 Boolean decision variables and 14,369 linear constraint equations."),
        ("Why Constraint Programming? ", "Unlike genetic algorithms or simulated annealing, CP-SAT uses SAT-based Conflict-Driven Clause Learning (CDCL) and integer programming cuts to mathematically prove feasibility and optimality."),
        ("Search Worker Parallelism: ", "Harnesses all available CPU cores (multi-threaded search workers) for parallel branch-and-bound search tree exploration.")
    ]
    for h_txt, b_txt in pts:
        bp = tf.add_paragraph()
        bp.text = "* " + h_txt + b_txt
        bp.font.name = FONT_BODY
        bp.font.size = Pt(10)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(7)

    # Col 2: Two-Phase Optimization Methodology
    add_card(slide, Inches(0.8) + col_w + gap, top_pos, col_w, card_h)
    tb = slide.shapes.add_textbox(Inches(1.05) + col_w + gap, top_pos + Inches(0.2), col_w - Inches(0.5), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Two-Phase Solving Strategy"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD
    p.space_after = Pt(8)
    
    pts = [
        ("Phase 1 - Hard Feasibility: ", "Solves pure satisfaction baseline: find x satisfying H1 and H2 and ... and H10. Completed in 3.9-4.7 seconds for all 26 classes."),
        ("Warm-Start Hint Injection: ", "Extracts Boolean variable assignments from Phase 1 solution and injects them as search hints (model.addHint(var, val)) into Phase 2."),
        ("Phase 2 - Soft Optimization: ", "Minimizes objective function: Minimize Sum(W_i * Penalty_i) for i in {S1..S6} while preserving all hard constraints."),
        ("Bounded Optimization Time: ", "Runs within user-configured timeout (default 60s, typically settling in 30-40s) returning the best solution found.")
    ]
    for h_txt, b_txt in pts:
        bp = tf.add_paragraph()
        bp.text = "* " + h_txt + b_txt
        bp.font.name = FONT_BODY
        bp.font.size = Pt(10)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(7)

def build_slide_8(prs):
    # Slide 8: Constraint Catalog: Hard Operational Rules (H1 to H10)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Constraint Catalog", "Enforced Hard Operational Rules (H1 to H10)",
               "Ten Strict Mathematical Rules Guaranteeing Zero Clashes and Operational Feasibility")
    
    col_w = Inches(5.65)
    gap_x = Inches(0.43)
    row_h = Inches(0.92)
    gap_y = Inches(0.12)
    top_pos = Inches(1.7)
    x0 = Inches(0.8)
    
    rules = [
        ("H1", "Class Teacher Period 1", "Class teacher assigned to Period 1 whenever teaching their assigned class."),
        ("H2", "Max 2 Continuous Lectures", "Faculty cannot teach >2 continuous periods without break; recess resets count."),
        ("H3", "Single Assignment / Slot", "Exactly one subject-teacher allocated per class time slot (48 weekly periods)."),
        ("H4", "Teacher Single-Booking", "A teacher can instruct at most one class division in any given period slot."),
        ("H5", "Computer Lab Exclusivity", "Across all 26 classes, at most one class may occupy the Computer Lab per slot."),
        ("H6", "P.E. Prohibited in Period 1", "Physical Education cannot be scheduled in morning Period 1 for any class."),
        ("H7", "P.E. Prohibited After Recess", "Physical Education prohibited in Period 5 immediately following morning recess."),
        ("H8", "Max 2 Activity Periods / Day", "A class division can have at most 2 co-curricular/activity periods per day."),
        ("H9", "Exact Allocation Fulfillment", "Every atomic (class, teacher, subject) period count K is exactly fulfilled."),
        ("H10", "P.E. & Comp Recess Partition", "When P.E. & Computer occur on same day, one must be before recess & one after.")
    ]
    
    for idx, (code, title, desc) in enumerate(rules):
        col = idx // 5
        row = idx % 5
        cx = x0 + col * (col_w + gap_x)
        cy = top_pos + row * (row_h + gap_y)
        
        add_card(slide, cx, cy, col_w, row_h)
        
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + Inches(0.15), cy + Inches(0.18), Inches(0.65), Inches(0.55))
        badge.fill.solid()
        badge.fill.fore_color.rgb = ACCENT_TERRACOTTA
        badge.line.fill.background()
        tf_b = badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = code
        p_b.font.name = FONT_HEADING
        p_b.font.size = Pt(11)
        p_b.font.bold = True
        p_b.font.color.rgb = TEXT_LIGHT
        p_b.alignment = PP_ALIGN.CENTER
        
        tb = slide.shapes.add_textbox(cx + Inches(0.9), cy + Inches(0.12), col_w - Inches(1.0), row_h - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = ACCENT_GOLD
        p0.space_after = Pt(1)
        
        p1 = tf.add_paragraph()
        p1.text = desc
        p1.font.name = FONT_BODY
        p1.font.size = Pt(9.2)
        p1.font.color.rgb = TEXT_LIGHT

def build_slide_9(prs):
    # Slide 9: Pedagogical Optimization: Soft Objectives (S1 to S6)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Optimization Objectives", "Pedagogical Quality Optimization (S1 to S6)",
               "Balancing Faculty Workloads and Enhancing Student Learning Experience")
    
    add_card(slide, Inches(0.8), Inches(1.7), Inches(11.733), Inches(0.75), border_color=ACCENT_GOLD)
    tb_f = slide.shapes.add_textbox(Inches(1.0), Inches(1.75), Inches(11.333), Inches(0.65))
    tf_f = tb_f.text_frame
    tf_f.word_wrap = True
    pf = tf_f.paragraphs[0]
    pf.text = "Objective Function:  Minimize Sum(Weight_i * Penalty_i)  for i in {S1, S2, S3, S4, S5, S6}"
    pf.font.name = FONT_HEADING
    pf.font.size = Pt(13)
    pf.font.bold = True
    pf.font.color.rgb = ACCENT_GOLD
    pf.alignment = PP_ALIGN.CENTER
    
    card_w = Inches(3.64)
    card_h = Inches(2.0)
    gap_x = Inches(0.4)
    gap_y = Inches(0.25)
    x0 = Inches(0.8)
    y0 = Inches(2.65)
    
    softs = [
        ("S1: Subject Spread Across Week", "Weight: 100", ACCENT_TERRACOTTA,
         "Distributes multi-period subjects evenly across weekdays. At most 1 lecture/day for subjects with K <= 6 periods/week, preventing subject fatigue."),
        ("S2: Non-Consecutive 2-Period Days", "Weight: 80", ACCENT_GOLD,
         "Subjects with exactly 2 lectures/week should not appear on adjacent days. Targets a gap >= 2 days to maintain learning retention."),
        ("S3: Repeated Lecture Spacing", "Weight: 50", ACCENT_CYAN,
         "Penalizes tight clustering of repeated subject sessions across the week. Maximizes uniform pedagogical pacing across syllabus units."),
        ("S4: Teacher Daily Workload Balance", "Weight: 40", ACCENT_EMERALD,
         "Minimizes variance in daily teaching load around the weekly average (weekly load / 6). Avoids heavy 6-lecture days alternating with 1-lecture days."),
        ("S5: Teacher Idle Gap Minimization", "Weight: 30", ACCENT_PURPLE,
         "Minimizes unassigned idle gap window periods between a teacher's first and last period of each day. Prevents unproductive waiting periods."),
        ("S6: Class Academic Variety", "Weight: 20", ACCENT_CYAN,
         "Prevents clustering more than 3 consecutive heavy academic lectures (Math, Science) without a co-curricular or linguistic break.")
    ]
    
    for idx, (title, weight, color, desc) in enumerate(softs):
        row = idx // 3
        col = idx % 3
        cx = x0 + col * (card_w + gap_x)
        cy = y0 + row * (card_h + gap_y)
        
        add_card(slide, cx, cy, card_w, card_h)
        tb = slide.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.15), card_w - Inches(0.4), card_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = color
        
        p1 = tf.add_paragraph()
        p1.text = weight
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_GOLD
        p1.space_after = Pt(4)
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(9.2)
        p2.font.color.rgb = TEXT_LIGHT

def build_slide_10(prs):
    # Slide 10: Modern User Interface & Studio Features
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "User Experience", "Command Center & Interactive Timetable Studio",
               "Modern Desktop User Interface Built with React 19, TypeScript, and SIES Visual Identity")
    
    col_w = Inches(3.64)
    gap = Inches(0.4)
    top_pos = Inches(1.7)
    card_h = Inches(4.3)
    
    # Col 1: Command Center
    add_card(slide, Inches(0.8), top_pos, col_w, card_h)
    tb = slide.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Command Center"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.space_after = Pt(8)
    
    pts = [
        ("System Health Dashboard: ", "Real-time indicators verifying loopback API daemon readiness, active port, and memory footprint."),
        ("Dataset Detection: ", "Auto-detects and validates presence of all 3 required input Excel files in user workspace."),
        ("One-Click Solver Execution: ", "Triggers asynchronous solver with live multi-stage progress listener."),
        ("Solve Metric Summary: ", "Reports feasibility duration, optimization duration, and constraint verification counters.")
    ]
    for h_txt, b_txt in pts:
        bp = tf.add_paragraph()
        bp.text = "* " + h_txt + b_txt
        bp.font.name = FONT_BODY
        bp.font.size = Pt(9.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(8)

    # Col 2: School Setup View
    add_card(slide, Inches(0.8) + col_w + gap, top_pos, col_w, card_h)
    tb = slide.shapes.add_textbox(Inches(1.0) + col_w + gap, top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "School Setup View"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD
    p.space_after = Pt(8)
    
    pts = [
        ("Configurable Academic Structure: ", "Dynamically configure standards (5 to 10) and division counts (A through E)."),
        ("Bell Schedule & Timings: ", "Customizable period counts (default 8/day), period start/end timings, and recess slot placement."),
        ("Faculty & Subject Directory: ", "Inspect and manage 44 teachers, subject categories (core, activity, special), and class teachers."),
        ("Model Validation: ", "Live validation of configuration consistency before initiating timetable generation.")
    ]
    for h_txt, b_txt in pts:
        bp = tf.add_paragraph()
        bp.text = "* " + h_txt + b_txt
        bp.font.name = FONT_BODY
        bp.font.size = Pt(9.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(8)

    # Col 3: Timetable Studio & Manual Editing
    add_card(slide, Inches(0.8) + (col_w + gap) * 2, top_pos, col_w, card_h)
    tb = slide.shapes.add_textbox(Inches(1.0) + (col_w + gap) * 2, top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Timetable Studio & Editing"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD
    p.space_after = Pt(8)
    
    pts = [
        ("Day * Period Grid Viewers: ", "Color-coded weekly grids for all 26 classes and 44 teachers with recess blocks."),
        ("In-Place Manual Slot Swapping: ", "Select any two periods to swap assignments directly on the visual grid."),
        ("Live Instant Constraint Audit: ", "Validates prospective swaps against H1 to H10 and S1 to S6 before applying changes."),
        ("Full Undo / Redo History: ", "Complete reversible editing stack allowing safe schedule experimentation.")
    ]
    for h_txt, b_txt in pts:
        bp = tf.add_paragraph()
        bp.text = "* " + h_txt + b_txt
        bp.font.name = FONT_BODY
        bp.font.size = Pt(9.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(8)

    # Bottom Banner
    add_card(slide, Inches(0.8), Inches(6.15), Inches(11.733), Inches(0.7), border_color=ACCENT_PURPLE, bg_color=CARD_BG_ALT)
    tb_b = slide.shapes.add_textbox(Inches(1.0), Inches(6.2), Inches(11.333), Inches(0.6))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    pb = tf_b.paragraphs[0]
    pb.text = "Interactive Workflow:  Command Center (Health & Launch)  ->  School Setup (Custom Rules)  ->  Timetable Studio (Grid View, Swap & Export)"
    pb.font.name = FONT_HEADING
    pb.font.size = Pt(10.5)
    pb.font.bold = True
    pb.font.color.rgb = ACCENT_PURPLE
    pb.alignment = PP_ALIGN.CENTER

def build_slide_11(prs):
    # Slide 11: Production Deliverables & Multi-Format Exporters
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Export Suite", "Production Deliverables & Multi-Format Exporters",
               "Automated High-Fidelity Publishing of Excel Workbooks, Class PDFs, and Audit Reports")
    
    col_w = Inches(3.64)
    gap = Inches(0.4)
    top_pos = Inches(1.7)
    card_h = Inches(4.3)
    
    # Col 1: Master Excel
    add_card(slide, Inches(0.8), top_pos, col_w, card_h)
    tb = slide.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Master Excel Workbook"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD
    p.space_after = Pt(2)
    
    p_sub = tf.add_paragraph()
    p_sub.text = "timetable.xlsx (Apache POI)"
    p_sub.font.name = FONT_HEADING
    p_sub.font.size = Pt(10)
    p_sub.font.bold = True
    p_sub.font.color.rgb = ACCENT_GOLD
    p_sub.space_after = Pt(8)
    
    pts = [
        ("Multi-Sheet Architecture: ", "Dedicated sheets for Summary metrics, All-Class Schedules, and All-Teacher Matrices."),
        ("Visual Styling & Recess: ", "Custom cell styles, styled headers, bold period boundaries, and color-coded morning recess blocks."),
        ("Workload Aggregations: ", "Computes total periods taught per faculty member and division-level weekly totals."),
        ("Administrative Ready: ", "Formatted for immediate printing or distribution to academic heads.")
    ]
    for h_txt, b_txt in pts:
        bp = tf.add_paragraph()
        bp.text = "* " + h_txt + b_txt
        bp.font.name = FONT_BODY
        bp.font.size = Pt(9.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(6)

    # Col 2: Vector PDFs
    add_card(slide, Inches(0.8) + col_w + gap, top_pos, col_w, card_h)
    tb = slide.shapes.add_textbox(Inches(1.0) + col_w + gap, top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Printable Per-Class PDFs"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TERRACOTTA
    p.space_after = Pt(2)
    
    p_sub = tf.add_paragraph()
    p_sub.text = "output/pdf/*.pdf (OpenPDF)"
    p_sub.font.name = FONT_HEADING
    p_sub.font.size = Pt(10)
    p_sub.font.bold = True
    p_sub.font.color.rgb = ACCENT_GOLD
    p_sub.space_after = Pt(8)
    
    pts = [
        ("26 Dedicated Class PDFs: ", "Individual PDF documents generated for each division (5-A through 10-D)."),
        ("High-Quality Vector Layout: ", "Landscape layout with institutional header, academic year, class teacher, and timings."),
        ("Legible Period Matrices: ", "Crisp table borders, bold subject names, and faculty code initials for easy reading."),
        ("Noticeboard Ready: ", "Standard A4 landscape size ready for immediate classroom display.")
    ]
    for h_txt, b_txt in pts:
        bp = tf.add_paragraph()
        bp.text = "* " + h_txt + b_txt
        bp.font.name = FONT_BODY
        bp.font.size = Pt(9.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(6)

    # Col 3: CSV Audit Suite
    add_card(slide, Inches(0.8) + (col_w + gap) * 2, top_pos, col_w, card_h)
    tb = slide.shapes.add_textbox(Inches(1.0) + (col_w + gap) * 2, top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Machine-Readable CSVs"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.space_after = Pt(2)
    
    p_sub = tf.add_paragraph()
    p_sub.text = "9 Audit & Data Reports"
    p_sub.font.name = FONT_HEADING
    p_sub.font.size = Pt(10)
    p_sub.font.bold = True
    p_sub.font.color.rgb = ACCENT_GOLD
    p_sub.space_after = Pt(8)
    
    pts = [
        ("student_timetable.csv: ", "Complete 26-class schedule matrix indexed by Class, Day, Period, Subject, Teacher."),
        ("teacher_timetable.csv: ", "Complete 44-teacher matrix enabling individual faculty schedule lookup."),
        ("reconciliation_discrepancies.csv: ", "Detailed log of upstream allotment mismatches between standard and teacher sheets."),
        ("validation_report.csv: ", "Granular audit report documenting compliance across all constraints.")
    ]
    for h_txt, b_txt in pts:
        bp = tf.add_paragraph()
        bp.text = "* " + h_txt + b_txt
        bp.font.name = FONT_BODY
        bp.font.size = Pt(9.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(6)

    # Bottom Banner
    add_card(slide, Inches(0.8), Inches(6.15), Inches(11.733), Inches(0.7), border_color=ACCENT_EMERALD, bg_color=CARD_BG_ALT)
    tb_b = slide.shapes.add_textbox(Inches(1.0), Inches(6.2), Inches(11.333), Inches(0.6))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    pb = tf_b.paragraphs[0]
    pb.text = "Local Storage Target:  %APPDATA%/com.sies.timetable.desktop/output/  *  Strictly local offline generation with zero cloud telemetry"
    pb.font.name = FONT_HEADING
    pb.font.size = Pt(10.5)
    pb.font.bold = True
    pb.font.color.rgb = ACCENT_EMERALD
    pb.alignment = PP_ALIGN.CENTER

def build_slide_12(prs):
    # Slide 12: Independent Validation & Engineering Verification
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Reliability & Verification", "Independent Validation & Automated Testing",
               "Zero-Trust Architecture and 100% Automated Test Suite Passing across All Components")
    
    col_w = Inches(5.65)
    gap = Inches(0.43)
    top_pos = Inches(1.7)
    card_h = Inches(3.3)
    
    # Col 1: Zero-Trust Validator
    add_card(slide, Inches(0.8), top_pos, col_w, card_h)
    tb = slide.shapes.add_textbox(Inches(1.05), top_pos + Inches(0.18), col_w - Inches(0.5), card_h - Inches(0.35))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Zero-Trust ScheduleValidator"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.space_after = Pt(6)
    
    pts = [
        ("Independent Post-Solve Audit: ", "ScheduleValidator.java inspects the candidate Schedule object completely from scratch with zero reliance on solver internal state."),
        ("Axiomatic Verification: ", "Re-evaluates every slot against all hard constraints H1 to H10 and soft objectives S1 to S6 independently."),
        ("Granular Coordinate Diagnostics: ", "Pinpoints exact failure coordinates (Day, Period, Class, Teacher, Subject) with human-readable remediation context."),
        ("Automated Quality Score: ", "Calculates objective penalty index quantifying timetable quality across faculty workload balance and subject pacing.")
    ]
    for h_txt, b_txt in pts:
        bp = tf.add_paragraph()
        bp.text = "* " + h_txt + b_txt
        bp.font.name = FONT_BODY
        bp.font.size = Pt(9.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(4)

    # Col 2: Verified Test Suite
    add_card(slide, Inches(0.8) + col_w + gap, top_pos, col_w, card_h)
    tb = slide.shapes.add_textbox(Inches(1.05) + col_w + gap, top_pos + Inches(0.18), col_w - Inches(0.5), card_h - Inches(0.35))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Verified Automated Test Suite"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD
    p.space_after = Pt(6)
    
    pts = [
        ("34 / 34 Java Tests Passing (0 Failures): ", "Full test suite executed via Maven Surefire in 3m 19s (DynamicStructure, EndToEnd, Exporters, Validation)."),
        ("14 REST API Contract Tests: ", "HttpServerIntegrationTest verifies all HTTP endpoints, preflight CORS, status codes, and shutdown signals."),
        ("Empirical Benchmark Suite: ", "SolverPerformanceBenchmarkTest validates solve times, memory footprint, and cancellation responsiveness (<32 ms)."),
        ("Frontend Linting & Bundling: ", "Oxlint verified (0 errors, 0 warnings across 21 files); production frontend compiles in 515 ms.")
    ]
    for h_txt, b_txt in pts:
        bp = tf.add_paragraph()
        bp.text = "* " + h_txt + b_txt
        bp.font.name = FONT_BODY
        bp.font.size = Pt(9.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(4)

    # Bottom Table: Test Suite Coverage Matrix
    tbl_top = Inches(5.15)
    tbl_h = Inches(1.7)
    add_card(slide, Inches(0.8), tbl_top, Inches(11.733), tbl_h, border_color=ACCENT_GOLD, bg_color=CARD_BG_ALT)
    
    tb_t = slide.shapes.add_textbox(Inches(1.0), tbl_top + Inches(0.12), Inches(11.333), tbl_h - Inches(0.24))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    
    pt0 = tf_t.paragraphs[0]
    pt0.text = "AUTOMATED TEST SUITE COVERAGE (34 TESTS ACROSS 9 SUITES - 100% PASS RATE)"
    pt0.font.name = FONT_HEADING
    pt0.font.size = Pt(10)
    pt0.font.bold = True
    pt0.font.color.rgb = ACCENT_GOLD
    pt0.space_after = Pt(4)
    
    matrix = [
        ("DynamicStructureSolverTest (2)", "EndToEndSolverTest (1)", "ExcelTimetableExporterTest (1)"),
        ("HttpServerIntegrationTest (14)", "NormalizationTest (2)", "ReconciliationTest (5)"),
        ("ScheduleValidatorTest (3)", "SolverPerformanceBenchmarkTest (3)", "TimetableServiceTest (3)")
    ]
    for r1, r2, r3 in matrix:
        pr = tf_t.add_paragraph()
        pr.text = f"* {r1:<34}  * {r2:<32}  * {r3}"
        pr.font.name = "Consolas"
        pr.font.size = Pt(9.2)
        pr.font.color.rgb = TEXT_LIGHT
        pr.space_after = Pt(2)

def build_slide_13(prs):
    # Slide 13: Desktop Packaging & Performance Benchmarks
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Deployment & Benchmarks", "Desktop Packaging & Empirical Performance",
               "Self-Contained Windows Distribution Coupled with Rigorous Benchmark Profiling")
    
    top_y = Inches(1.7)
    card_w = Inches(2.7)
    card_h = Inches(1.3)
    gap_x = Inches(0.31)
    x0 = Inches(0.8)
    
    metrics = [
        ("INSTALLER SIZE", "58.52 MB", ACCENT_CYAN, "Single-file NSIS setup"),
        ("SOLVE TIME (26 CLS)", "34 - 48s", ACCENT_EMERALD, "0 hard violations"),
        ("AUTOMATED TESTS", "34 / 34", ACCENT_GOLD, "100% pass rate (0 errors)"),
        ("DESKTOP RAM", "27.5 MB", ACCENT_PURPLE, "Tauri WebView2 shell")
    ]
    
    for idx, (label, val, color, note) in enumerate(metrics):
        cx = x0 + idx * (card_w + gap_x)
        add_card(slide, cx, top_y, card_w, card_h)
        tb = slide.shapes.add_textbox(cx + Inches(0.15), top_y + Inches(0.12), card_w - Inches(0.3), card_h - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = label
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(9)
        p0.font.bold = True
        p0.font.color.rgb = TEXT_MUTED
        
        p1 = tf.add_paragraph()
        p1.text = val
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(20)
        p1.font.bold = True
        p1.font.color.rgb = color
        
        p2 = tf.add_paragraph()
        p2.text = note
        p2.font.name = FONT_BODY
        p2.font.size = Pt(9)
        p2.font.color.rgb = TEXT_LIGHT

    bot_y = Inches(3.25)
    bot_h = Inches(3.55)
    b_col_w = Inches(5.65)
    b_gap = Inches(0.43)
    
    # Left: Turnkey Windows Packaging
    add_card(slide, Inches(0.8), bot_y, b_col_w, bot_h)
    tb = slide.shapes.add_textbox(Inches(1.05), bot_y + Inches(0.2), b_col_w - Inches(0.5), bot_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Turnkey Windows Packaging Strategy"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.space_after = Pt(6)
    
    pts = [
        ("Private Bundled JRE (48 MB): ", "Uses jlink with 11 core modules including jdk.zipfs (required by OR-Tools C++ DLL extractor) and jdk.httpserver."),
        ("No Admin Privileges Required: ", "NSIS configured with installMode: currentUser, installing cleanly into %LOCALAPPDATA% without UAC elevation."),
        ("Zero Pre-requisites: ", "End users never need to install Java, Maven, Python, or Node.js. Launches directly from Start Menu."),
        ("Data Isolation: ", "User configs and exports stored safely in %APPDATA%/com.sies.timetable.desktop/output.")
    ]
    for h_txt, b_txt in pts:
        bp = tf.add_paragraph()
        bp.text = "* " + h_txt + b_txt
        bp.font.name = FONT_BODY
        bp.font.size = Pt(9.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(5)

    # Right: Empirical Solver Benchmarks
    add_card(slide, Inches(0.8) + b_col_w + b_gap, bot_y, b_col_w, bot_h)
    tb = slide.shapes.add_textbox(Inches(1.05) + b_col_w + b_gap, bot_y + Inches(0.2), b_col_w - Inches(0.5), bot_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Empirical Solver Benchmarks"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD
    p.space_after = Pt(6)
    
    pts = [
        ("Small Dataset (4 Classes): ", "1,232 variables, 1,381 constraints. Phase 1 Feasible: 0.15s; Phase 2 Optimized: 1.62s; Total: 1.79s."),
        ("Production Dataset (26 Classes): ", "18,864 variables, 14,369 constraints. Phase 1 Feasible: 3.92s; Phase 2 Optimized: 29.90s; Total: 34.09s."),
        ("Hard Violations: ", "Exactly 0 hard constraint violations across all generated test and production timetables."),
        ("Cancellation Responsiveness: ", "Active solve job aborts cleanly in <32 ms upon cancellation request.")
    ]
    for h_txt, b_txt in pts:
        bp = tf.add_paragraph()
        bp.text = "* " + h_txt + b_txt
        bp.font.name = FONT_BODY
        bp.font.size = Pt(9.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(5)

def build_slide_14(prs):
    # Slide 14: Limitations, Roadmap & Conclusion
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Summary & Future Outlook", "System Limitations, Future Roadmap & Conclusion",
               "Reflecting on Operational Capabilities and Academic Engineering Impact")
    
    col_w = Inches(3.64)
    gap = Inches(0.4)
    top_pos = Inches(1.7)
    card_h = Inches(4.3)
    
    # Col 1: Current Limitations
    add_card(slide, Inches(0.8), top_pos, col_w, card_h)
    tb = slide.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Current Limitations"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TERRACOTTA
    p.space_after = Pt(8)
    
    pts = [
        ("Desktop-Only Deployment: ", "Architecture runs locally on loopback; does not currently offer a multi-tenant web cloud portal."),
        ("Single Active Solver Job: ", "Local API queues or rejects simultaneous concurrent solves to preserve CPU solver threads."),
        ("Unsigned Binary: ", "Lacks a commercial EV code signing certificate; Windows SmartScreen warning occurs on first launch."),
        ("Single Lab Dimension: ", "Currently enforces exclusivity for Computer Lab; specialized science labs modeled as standard classrooms.")
    ]
    for h_txt, b_txt in pts:
        bp = tf.add_paragraph()
        bp.text = "* " + h_txt + b_txt
        bp.font.name = FONT_BODY
        bp.font.size = Pt(9.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(8)

    # Col 2: Future Roadmap
    add_card(slide, Inches(0.8) + col_w + gap, top_pos, col_w, card_h)
    tb = slide.shapes.add_textbox(Inches(1.0) + col_w + gap, top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Future Roadmap"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD
    p.space_after = Pt(8)
    
    pts = [
        ("Multi-Facility Room Scheduling: ", "Extend spatial resource scheduling to Science Labs, AV Rooms, and Library sessions."),
        ("Teacher Absence Substitute Engine: ", "Daily dynamic re-optimizer to allocate free teachers when faculty report sudden sick leave."),
        ("Multi-School Cloud Sync: ", "Optional central repository sync for school group administrators managing multiple campuses."),
        ("Commercial Code Signing: ", "Integration with Microsoft Authenticode EV signing for zero-warning enterprise deployment.")
    ]
    for h_txt, b_txt in pts:
        bp = tf.add_paragraph()
        bp.text = "* " + h_txt + b_txt
        bp.font.name = FONT_BODY
        bp.font.size = Pt(9.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(8)

    # Col 3: Conclusion
    add_card(slide, Inches(0.8) + (col_w + gap) * 2, top_pos, col_w, card_h, border_color=ACCENT_EMERALD)
    tb = slide.shapes.add_textbox(Inches(1.0) + (col_w + gap) * 2, top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Project Conclusion"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD
    p.space_after = Pt(8)
    
    pts = [
        ("Engineering Success: ", "Engineered an end-to-end, mathematically verified timetabling solution that replaces 4 weeks of manual toil with a 40-second solve."),
        ("Zero-Violation Reliability: ", "Achieved 100% adherence to all 10 hard constraints and verified compliance via independent zero-trust validation."),
        ("Production Turnkey Usability: ", "Delivered a lightweight, self-contained desktop installer ready for school administrators."),
        ("Open for Questions: ", "Thank you! Open to faculty and student technical questions.")
    ]
    for h_txt, b_txt in pts:
        bp = tf.add_paragraph()
        bp.text = "* " + h_txt + b_txt
        bp.font.name = FONT_BODY
        bp.font.size = Pt(9.5)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_after = Pt(8)

    # Bottom Banner
    add_card(slide, Inches(0.8), Inches(6.15), Inches(11.733), Inches(0.7), border_color=ACCENT_GOLD, bg_color=CARD_BG_ALT)
    tb_b = slide.shapes.add_textbox(Inches(1.0), Inches(6.2), Inches(11.333), Inches(0.6))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    pb = tf_b.paragraphs[0]
    pb.text = "Key Takeaway: Autonomous constraint satisfaction eliminates human scheduling errors, ensures faculty equity, and guarantees 100% policy compliance."
    pb.font.name = FONT_HEADING
    pb.font.size = Pt(10.5)
    pb.font.bold = True
    pb.font.color.rgb = ACCENT_GOLD
    pb.alignment = PP_ALIGN.CENTER

def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    print("Building Slide 1: Title...")
    build_slide_1(prs)
    print("Building Slide 2: Problem Statement...")
    build_slide_2(prs)
    print("Building Slide 3: Project Objectives...")
    build_slide_3(prs)
    print("Building Slide 4: Technology Stack...")
    build_slide_4(prs)
    print("Building Slide 5: System Architecture...")
    build_slide_5(prs)
    print("Building Slide 6: Timetable Generation Workflow...")
    build_slide_6(prs)
    print("Building Slide 7: OR-Tools CP-SAT Formulation...")
    build_slide_7(prs)
    print("Building Slide 8: Hard Constraints Catalog...")
    build_slide_8(prs)
    print("Building Slide 9: Soft Objectives Optimization...")
    build_slide_9(prs)
    print("Building Slide 10: User Experience & Studio...")
    build_slide_10(prs)
    print("Building Slide 11: Production Deliverables...")
    build_slide_11(prs)
    print("Building Slide 12: Independent Validation & Testing...")
    build_slide_12(prs)
    print("Building Slide 13: Desktop Packaging & Benchmarks...")
    build_slide_13(prs)
    print("Building Slide 14: Limitations, Roadmap & Conclusion...")
    build_slide_14(prs)
    
    out_dir = os.path.join(os.getcwd(), "presentation")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "SIES_Timetable_Generator_Presentation.pptx")
    prs.save(out_path)
    print(f"\n[SUCCESS] Presentation saved to: {out_path} ({len(prs.slides)} slides)")

if __name__ == "__main__":
    main()
